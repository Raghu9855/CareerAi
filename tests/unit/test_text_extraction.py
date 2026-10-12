"""Unit tests for text extraction pipeline and edge cases."""

import io

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.core.errors import ValidationError
from app.services.text_extraction import (
    clean_extracted_text,
    extract_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_from_txt,
)


def test_clean_extracted_text() -> None:
    """Verify cleaning collapses whitespace, handles CRLF, and strips edges."""
    assert clean_extracted_text("") == ""
    raw = "  Hello \r\n\r\n World  \t with   spaces \n\n\n\n Extra lines  "
    cleaned = clean_extracted_text(raw)
    assert "Hello" in cleaned
    assert "World with spaces" in cleaned
    assert "\n\n\n" not in cleaned


def test_extract_text_from_valid_pdf() -> None:
    """Verify extracting text from generated valid PDF."""
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    result = extract_text_from_pdf(pdf_bytes)
    assert isinstance(result, str)


def test_extract_text_from_corrupt_pdf_raises_validation_error() -> None:
    """Verify corrupt PDF bytes trigger ValidationError."""
    with pytest.raises(ValidationError) as exc:
        extract_text_from_pdf(b"This is definitely not a real PDF file %PDF")
    assert "Corrupt or unreadable PDF" in exc.value.message


def test_extract_text_from_valid_docx() -> None:
    """Verify extracting text from generated DOCX document."""
    doc = DocxDocument()
    doc.add_paragraph("Paragraph 1: Skills in FastAPI and Python.")
    doc.add_paragraph("Paragraph 2: Database architecture in MySQL.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extracted = extract_text_from_docx(docx_bytes)
    assert "Paragraph 1: Skills in FastAPI and Python." in extracted
    assert "Paragraph 2: Database architecture in MySQL." in extracted


def test_extract_text_from_corrupt_docx_raises_validation_error() -> None:
    """Verify invalid DOCX payload triggers ValidationError."""
    with pytest.raises(ValidationError) as exc:
        extract_text_from_docx(b"PK\x03\x04corrupt_docx_stream")
    assert "Corrupt or unreadable DOCX" in exc.value.message


def test_extract_text_from_txt_encodings() -> None:
    """Verify TXT extractor handles UTF-8 and latin-1 encodings."""
    utf8_bytes = "Experienced engineer with résumé and façade".encode()
    assert "résumé" in extract_text_from_txt(utf8_bytes)

    latin1_bytes = "Resume with special accents: café".encode("latin-1")
    assert "café" in extract_text_from_txt(latin1_bytes)


def test_extract_text_router_and_unsupported_format() -> None:
    """Verify extract_text router maps extensions and rejects unsupported types."""
    txt_bytes = b"Plain test content"
    assert extract_text(txt_bytes, ".txt") == "Plain test content"
    assert extract_text(txt_bytes, " .TXT ") == "Plain test content"

    with pytest.raises(ValidationError) as exc:
        extract_text(txt_bytes, ".exe")
    assert "Unsupported format for text extraction" in exc.value.message
